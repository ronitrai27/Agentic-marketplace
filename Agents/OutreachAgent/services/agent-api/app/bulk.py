"""Excel workbook parsing and email templating for bulk outreach uploads.

The first non-empty row of an uploaded workbook defines the available fields.
Column headers are matched flexibly (``Email``, ``E-Mail``, ``Author Name``,
``Company``, ...) and every column becomes a ``{{token}}`` that can be used in
the email subject/body templates. One personalized email is rendered per
recipient row; rows without a special value fall back cleanly instead of
leaving placeholders behind.
"""

import io
import re
from dataclasses import dataclass, field
from typing import Mapping

from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter, ValidationError

from .models import Recipient

MAX_BULK_RECIPIENTS = 50
MAX_BULK_FILE_BYTES = 2 * 1024 * 1024
MAX_SUBJECT_TEMPLATE_CHARS = 500
MAX_BODY_TEMPLATE_CHARS = 10_000
ALLOWED_EXTENSIONS = (".xlsx", ".xlsm")

# Canonical roles and the header spellings that map onto them. Matching is
# done on a normalized header key (lowercase, punctuation collapsed), so
# "E-Mail", "e mail" and "EMAIL" all resolve to the same role.
ROLE_ALIASES: dict[str, set[str]] = {
    "email": {"email", "email address", "emailid", "email id", "mail", "mail address", "e mail"},
    "name": {
        "name", "full name", "fullname", "author name", "authorname", "author",
        "recipient name", "contact name", "contactname", "person", "display name", "contact",
    },
    "first_name": {"first name", "firstname", "given name", "forename"},
    "last_name": {"last name", "lastname", "surname", "family name"},
    "company": {"company", "company name", "organisation", "organization", "org", "brand", "firm", "business"},
    "role": {"role", "title", "designation", "job title", "jobtitle", "position"},
    "context": {"context", "notes", "note", "background", "details", "remarks"},
}

_EMAIL = TypeAdapter(EmailStr)

# Compact (space-free) alias sets so "E-Mail Address" -> "e mail address" ->
# "emailaddress" still matches the "email address" alias.
_ROLE_COMPACT: dict[str, set[str]] = {
    role: {alias.replace(" ", "") for alias in aliases}
    for role, aliases in ROLE_ALIASES.items()
}


class BulkParseError(ValueError):
    """The uploaded workbook cannot be used for bulk outreach."""


@dataclass(frozen=True)
class DetectedColumn:
    """A header cell from row 1 and the token it exposes to templates."""

    token: str
    column: str
    role: str  # email | name | first_name | last_name | company | role | context | extra


@dataclass(frozen=True)
class SkippedRow:
    row: int
    reason: str


@dataclass(frozen=True)
class ParsedRow:
    row: int
    recipient: Recipient
    fields: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class BulkParseResult:
    rows: list[ParsedRow]
    skipped: list[SkippedRow]
    total_rows: int
    columns: list[DetectedColumn]

    @property
    def recipients(self) -> list[Recipient]:
        return [parsed.recipient for parsed in self.rows]


def _key(value: object) -> str:
    """Normalize a header/value into a stable token key."""
    text = re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()
    return re.sub(r"\s+", " ", text)


def _token(key: str) -> str:
    return key.replace(" ", "_")


def _cell(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _detect_columns(header_keys: list[str], header_cells: list[str]) -> list[DetectedColumn]:
    """Map row-1 headers onto canonical roles; unmatched headers stay as extras.

    The result is index-aligned with the header row so callers can use the
    position to read data cells safely.
    """
    columns: list[DetectedColumn] = []
    used_tokens: set[str] = set()
    role_taken: set[str] = set()
    for index, (key, cell) in enumerate(zip(header_keys, header_cells)):
        role = "extra"
        compact_key = key.replace(" ", "")
        for canonical, aliases in ROLE_ALIASES.items():
            matched = key in aliases or compact_key in _ROLE_COMPACT[canonical]
            if matched and canonical not in role_taken:
                role = canonical
                break
        token = role if role != "extra" else _token(key)
        if not token:
            token = f"column_{index + 1}"
        if token in used_tokens:
            suffix = 2
            while f"{token}_{suffix}" in used_tokens:
                suffix += 1
            token = f"{token}_{suffix}"
        used_tokens.add(token)
        if role != "extra":
            role_taken.add(role)
        columns.append(DetectedColumn(token=token, column=cell or f"Column {index + 1}", role=role))
    return columns


def _name_from_email(email: str) -> str:
    """Derive a human-looking name from an email local part as a fallback."""
    local = email.split("@", 1)[0]
    parts = re.split(r"[._\-+0-9]+", local)
    words = [part for part in parts if part.isalpha()][:2]
    if not words:
        cleaned = re.sub(r"[^A-Za-z ]", "", local).strip()
        return cleaned.title()
    return " ".join(word.capitalize() for word in words)


def parse_bulk_workbook(filename: str | None, data: bytes, locale: str = "en") -> BulkParseResult:
    """Parse an uploaded .xlsx workbook into recipients with template fields."""
    if not filename or not filename.lower().endswith(ALLOWED_EXTENSIONS):
        raise BulkParseError("Upload a .xlsx Excel file with an email column in row 1.")
    if not data:
        raise BulkParseError("The uploaded file is empty.")

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise BulkParseError("The file could not be read as an Excel workbook.") from error

    try:
        worksheet = workbook.active
        if worksheet is None:
            raise BulkParseError("The workbook has no worksheet.")
        header_cells: list[str] | None = None
        data_rows: list[tuple[int, list[object]]] = []
        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            cells = list(values)
            if header_cells is None:
                if not any(_cell(cell) for cell in cells):
                    continue
                header_cells = [_cell(cell) for cell in cells]
                continue
            if not any(_cell(cell) for cell in cells):
                continue
            data_rows.append((row_number, cells))
    finally:
        workbook.close()

    if header_cells is None:
        raise BulkParseError("The workbook is empty. Add a header row with an 'email' column.")
    if not any(header_cells):
        raise BulkParseError("The workbook has no headers. Add a header row with an 'email' column.")

    header_keys = [_key(cell) for cell in header_cells]
    columns = _detect_columns(header_keys, header_cells)
    if not any(column.role == "email" for column in columns):
        raise BulkParseError(
            "Missing 'email' column in row 1. Use a header like 'email', 'email address', or 'mail'."
        )
    if not data_rows:
        raise BulkParseError("No recipient rows found. Add rows below the header row.")
    if len(data_rows) > MAX_BULK_RECIPIENTS:
        raise BulkParseError(
            f"Bulk sends are limited to {MAX_BULK_RECIPIENTS} recipients; this file has {len(data_rows)} rows."
        )

    email_index = next(index for index, column in enumerate(columns) if column.role == "email")
    name_index = next((index for index, column in enumerate(columns) if column.role == "name"), None)
    first_index = next((index for index, column in enumerate(columns) if column.role == "first_name"), None)
    last_index = next((index for index, column in enumerate(columns) if column.role == "last_name"), None)
    context_index = next((index for index, column in enumerate(columns) if column.role == "context"), None)

    rows: list[ParsedRow] = []
    skipped: list[SkippedRow] = []
    for row_number, cells in data_rows:
        def cell_at(index: int | None) -> str:
            if index is None or index >= len(cells):
                return ""
            return _cell(cells[index])

        email = cell_at(email_index)
        if not email:
            skipped.append(SkippedRow(row=row_number, reason="Missing email address."))
            continue
        try:
            valid_email = str(_EMAIL.validate_python(email))
        except ValidationError:
            skipped.append(SkippedRow(row=row_number, reason=f"'{email}' is not a valid email address."))
            continue

        # Row fields: every column (canonical token + original header token).
        fields: dict[str, str] = {}
        for index, column in enumerate(columns):
            value = cell_at(index)
            fields[column.token] = value
            original_key = _token(header_keys[index]) if index < len(header_keys) else ""
            if original_key:
                fields.setdefault(original_key, value)

        name = cell_at(name_index)
        if not name and first_index is not None:
            name = " ".join(part for part in (cell_at(first_index), cell_at(last_index)) if part).strip()
        if not name:
            name = _name_from_email(valid_email)
        fields["name"] = name
        fields.setdefault("email", valid_email)
        first_name = cell_at(first_index) if first_index is not None else ""
        if not first_name and name:
            first_name = name.split()[0]
        fields["first_name"] = first_name

        if not name:
            skipped.append(SkippedRow(row=row_number, reason="Missing name and no usable email name."))
            continue

        rows.append(
            ParsedRow(
                row=row_number,
                recipient=Recipient(
                    recipient_id=f"bulk-row-{row_number}",
                    name=name,
                    email=valid_email,
                    locale=locale,
                    attributes={
                        "consent": True,
                        "supplied_context": cell_at(context_index),
                        "fields": dict(fields),
                    },
                ),
                fields=fields,
            )
        )

    if not rows:
        raise BulkParseError("No valid recipients found. Each row needs a valid email address.")
    return BulkParseResult(rows=rows, skipped=skipped, total_rows=len(data_rows), columns=columns)


def _template_values(fields: Mapping[str, str], context: Mapping[str, str]) -> dict[str, str]:
    """Merge row fields with run context and apply smart fallbacks."""
    values: dict[str, str] = {key: _cell(value) for key, value in context.items()}
    values.update({key: _cell(value) for key, value in fields.items()})

    email = values.get("email", "")
    name = values.get("name", "") or _name_from_email(email)
    values["name"] = name
    values["first_name"] = values.get("first_name", "") or (name.split()[0] if name else "")
    values["sender_signature"] = values.get("sender_signature") or values.get("sender_name") or ""
    return values


def render_template(template: str, values: Mapping[str, str]) -> str:
    """Substitute ``{{token}}`` placeholders and tidy the result.

    Unknown or empty placeholders are removed and the whitespace is repaired
    so no broken ``Hi ,`` style text or leftover ``{{...}}`` ever ships.
    """
    if not template:
        return ""

    def replace(match: re.Match[str]) -> str:
        token = _token(_key(match.group(1)))
        return values.get(token, "")

    rendered = re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", replace, template)
    rendered = re.sub(r"\{\{[^{}]*\}\}", "", rendered)
    rendered = re.sub(r"[ \t]+", " ", rendered)
    rendered = re.sub(r"\s+([,.;:!?)])", r"\1", rendered)
    rendered = re.sub(r"([(]\s+)", r"(", rendered)
    rendered = re.sub(r"\n{3,}", "\n\n", rendered)
    lines = [line.strip() for line in rendered.split("\n")]
    return "\n".join(lines).strip()


def render_bulk_email(
    subject_template: str,
    body_template: str,
    fields: Mapping[str, str],
    context: Mapping[str, str],
) -> tuple[str, str]:
    """Render one recipient's subject and body; empty when no body template."""
    if not body_template.strip():
        return "", ""
    values = _template_values(fields, context)
    subject = render_template(subject_template[:MAX_SUBJECT_TEMPLATE_CHARS], values)
    if not subject:
        subject = f"Note from {values.get('sender_name')}".strip() if values.get("sender_name") else "Introduction"
    body = render_template(body_template[:MAX_BODY_TEMPLATE_CHARS], values)
    if not body:
        return "", ""
    return subject, body
